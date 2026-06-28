from app.domain.knowledge import ParsedDocument, TextChunk


class DocumentParser:
    """Extract plain text from supported file formats."""

    def parse(self, filename: str, file_path: str, data: bytes) -> ParsedDocument:
        from app.domain.knowledge import compute_content_hash

        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        mime = self._mime_for_ext(ext)

        if ext == "pdf":
            content = self._parse_pdf(data)
        elif ext == "docx":
            content = self._parse_docx(data)
        else:
            content = data.decode("utf-8", errors="replace")

        content = content.strip()
        return ParsedDocument(
            filename=filename,
            file_path=file_path,
            mime_type=mime,
            content=content,
            content_hash=compute_content_hash(content),
            file_size=len(data),
            metadata={"extension": ext},
        )

    @staticmethod
    def _parse_pdf(data: bytes) -> str:
        from io import BytesIO

        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)

    @staticmethod
    def _parse_docx(data: bytes) -> str:
        from io import BytesIO

        from docx import Document

        doc = Document(BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    @staticmethod
    def _mime_for_ext(ext: str) -> str:
        mapping = {
            "pdf": "application/pdf",
            "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "md": "text/markdown",
            "markdown": "text/markdown",
            "txt": "text/plain",
            "py": "text/x-python",
            "ts": "text/typescript",
            "tsx": "text/typescript",
            "js": "text/javascript",
            "json": "application/json",
        }
        return mapping.get(ext, "text/plain")
